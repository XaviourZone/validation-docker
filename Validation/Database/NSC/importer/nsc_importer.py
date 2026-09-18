"""
NSC Importer - Loads NSC East and West (xlsx/csv) into nsc_vessels table.
"""

import argparse
import csv
import logging
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

# Add project root to path for imports if run directly
import sys
project_root = Path(__file__).resolve().parent.parent.parent.parent.parent
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from Validation.Database.common.database_utils import (
    open_database,
    create_import_tracking_tables,
    integrity_check,
)
from Validation.Database.common.file_utils import discover_files, sha256_file
from Validation.Database.common.logging_utils import setup_logger

import yaml

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def load_config(config_path):
    if not config_path:
        config_path = Path(__file__).resolve().parent.parent.parent / "config" / "database.yaml"
    
    with open(config_path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)

def clean_column_name(col):
    """Clean header column to be a valid SQLite column name."""
    c = str(col).strip()
    if not c or c.lower() == "none":
        return None  # Will be skipped
    chars_to_replace = [" ", "-", ".", "/", "\\", "(", ")", "[", "]"]
    for char in chars_to_replace:
        c = c.replace(char, "_")
    return c.upper()

def ensure_nsc_table(conn, headers):
    """Ensure nsc_vessels table exists and has all required columns."""
    cur = conn.cursor()
    
    # Check if table exists
    cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='nsc_vessels'")
    table_exists = cur.fetchone() is not None
    
    if not table_exists:
        cur.execute("CREATE TABLE nsc_vessels (_id INTEGER PRIMARY KEY AUTOINCREMENT, SOURCE_REGION TEXT)")
        
    cur.execute("PRAGMA table_info(nsc_vessels)")
    existing_columns = {row['name'] for row in cur.fetchall()}
    
    valid_headers = []
    for h in headers:
        cleaned = clean_column_name(h)
        if cleaned:
            valid_headers.append(cleaned)
            if cleaned not in existing_columns:
                cur.execute(f"ALTER TABLE nsc_vessels ADD COLUMN {cleaned} TEXT")
                existing_columns.add(cleaned)
        else:
            valid_headers.append(None) # Keep placeholder for skipped columns
            
    # Create indexes for lookup
    if not table_exists:
        if "ID_IMO" in existing_columns:
            cur.execute('CREATE INDEX idx_nsc_imo ON nsc_vessels("ID_IMO")')
        if "ID_MMSI" in existing_columns:
            cur.execute('CREATE INDEX idx_nsc_mmsi ON nsc_vessels("ID_MMSI")')
        cur.execute('CREATE INDEX idx_nsc_region ON nsc_vessels("SOURCE_REGION")')
        
    return valid_headers

def process_file(conn, file_path, batch_id, region, batch_size=5000):
    log = logging.getLogger("nsc_importer")
    start_time = time.monotonic()
    
    file_size = file_path.stat().st_size
    file_hash = sha256_file(file_path)
    
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO import_file 
        (batch_id, source_system, source_region, file_name, file_path, file_hash_sha256, file_size_bytes, target_table, discovered_at, status)
        VALUES (?, 'NSC', ?, ?, ?, ?, ?, 'nsc_vessels', datetime('now'), 'RUNNING')
    """, (batch_id, region, file_path.name, str(file_path), file_hash, file_size))
    file_id = cur.lastrowid
    conn.commit()

    rows_loaded = 0
    try:
        is_xlsx = file_path.suffix.lower() == '.xlsx'
        
        if is_xlsx:
            if not HAS_OPENPYXL:
                raise ImportError("openpyxl is not installed. Cannot read xlsx files. Use .csv instead.")
            
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            row_iter = ws.iter_rows(values_only=True)
            
            try:
                raw_headers = next(row_iter)
            except StopIteration:
                raise ValueError("Empty Excel file")
                
        else:
            f = open(file_path, "r", encoding="utf-8-sig")
            reader = csv.reader(f)
            try:
                raw_headers = next(reader)
            except StopIteration:
                f.close()
                raise ValueError("Empty CSV file")
            row_iter = reader

        valid_headers = ensure_nsc_table(conn, raw_headers)
        
        cols = ["SOURCE_REGION"] + [h for h in valid_headers if h]
        placeholders = ",".join(["?"] * len(cols))
        col_names = ",".join(cols)
        insert_sql = f'INSERT INTO nsc_vessels ({col_names}) VALUES ({placeholders})'
        
        batch = []
        for row in row_iter:
            # Filter out skipped columns
            filtered_row = [region]
            for i, val in enumerate(row):
                if i < len(valid_headers) and valid_headers[i]:
                    filtered_row.append(str(val) if val is not None else "")
            
            # If row had fewer columns than headers, pad
            while len(filtered_row) < len(cols):
                filtered_row.append("")
                
            batch.append(filtered_row)
            
            if len(batch) >= batch_size:
                conn.executemany(insert_sql, batch)
                rows_loaded += len(batch)
                batch = []
                
        if batch:
            conn.executemany(insert_sql, batch)
            rows_loaded += len(batch)
            
        if not is_xlsx:
            f.close()
        else:
            wb.close()
            
        conn.commit()
        
        cur.execute("""
            UPDATE import_file 
            SET loaded_at = datetime('now'), status = 'COMPLETED', rows_loaded = ?
            WHERE file_id = ?
        """, (rows_loaded, file_id))
        conn.commit()
        
        log.info(f"Loaded {file_path.name} ({region}): {rows_loaded} rows in {time.monotonic() - start_time:.2f}s")
        return True, rows_loaded
        
    except Exception as e:
        conn.rollback()
        log.error(f"Failed to process {file_path.name}: {e}")
        cur.execute("""
            UPDATE import_file 
            SET loaded_at = datetime('now'), status = 'FAILED', error_message = ?
            WHERE file_id = ?
        """, (str(e), file_id))
        conn.commit()
        return False, 0

def run_import(config, dry_run=False):
    log = logging.getLogger("nsc_importer")
    
    db_config = config.get('database', {}).get('nsc', {})
    import_config = config.get('imports', {}).get('nsc', {})
    
    db_path = Path(project_root) / db_config.get('path', 'Validation/Database/NSC/nsc.db')
    staging_path = Path(project_root) / import_config.get('staging_db', 'Validation/Database/NSC/nsc_staging.db')
    input_dir = Path(project_root) / import_config.get('input_dir', 'Validation/Database/NSC/RAW_DATA')
    batch_size = import_config.get('batch_size', 5000)
    
    east_dir = input_dir / "EAST"
    west_dir = input_dir / "WEST"
    
    # Fallback for testing with existing sample data
    fallback_dir = input_dir.parent.parent.parent.parent / "Sample data" / "NSC EAST and WEST"
    
    east_files = discover_files(east_dir, [".xlsx", ".csv"])
    west_files = discover_files(west_dir, [".xlsx", ".csv"])
    
    if not east_files and not west_files and fallback_dir.exists():
        log.info(f"Using fallback directory {fallback_dir}")
        for f in discover_files(fallback_dir, [".xlsx", ".csv"]):
            if "EAST" in f.name.upper():
                east_files.append(f)
            elif "WEST" in f.name.upper():
                west_files.append(f)
                
    all_files = east_files + west_files
    
    if not all_files:
        log.error("No NSC East/West files found.")
        return False
        
    if dry_run:
        log.info(f"DRY RUN: Discovered {len(east_files)} EAST files and {len(west_files)} WEST files.")
        for f in east_files: log.info(f"Would process (EAST): {f.name}")
        for f in west_files: log.info(f"Would process (WEST): {f.name}")
        return True
        
    log.info(f"Discovered {len(east_files)} EAST files and {len(west_files)} WEST files.")
    
    if staging_path.exists():
        staging_path.unlink()
        
    conn = open_database(staging_path)
    create_import_tracking_tables(conn)
    
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO import_batch (source_system, started_at, status, files_discovered)
        VALUES ('NSC', datetime('now'), 'RUNNING', ?)
    """, (len(all_files),))
    batch_id = cur.lastrowid
    conn.commit()
    
    total_loaded = 0
    files_loaded = 0
    errors = 0
    
    for f in east_files:
        success, rows = process_file(conn, f, batch_id, "EAST", batch_size=batch_size)
        if success:
            files_loaded += 1
            total_loaded += rows
        else:
            errors += 1
            
    for f in west_files:
        success, rows = process_file(conn, f, batch_id, "WEST", batch_size=batch_size)
        if success:
            files_loaded += 1
            total_loaded += rows
        else:
            errors += 1
            
    if errors > 0:
        log.error(f"Import failed with {errors} file errors. Staging database will NOT replace active.")
        status = 'FAILED'
        replace_db = False
    elif not integrity_check(conn):
        log.error("Staging database failed SQLite integrity check. Staging database will NOT replace active.")
        status = 'FAILED_INTEGRITY'
        replace_db = False
    else:
        status = 'COMPLETED'
        replace_db = True
        
    cur.execute("""
        UPDATE import_batch
        SET completed_at = datetime('now'), status = ?, files_loaded = ?, rows_loaded = ?, error_count = ?
        WHERE batch_id = ?
    """, (status, files_loaded, total_loaded, errors, batch_id))
    conn.commit()
    conn.close()
    
    if replace_db:
        log.info("Import successful. Replacing active database...")
        db_path.parent.mkdir(parents=True, exist_ok=True)
        
        if db_path.exists():
            backup_dir = db_path.parent.parent / "backups"
            backup_dir.mkdir(parents=True, exist_ok=True)
            backup_path = backup_dir / f"nsc_pre_replace_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
            shutil.copy2(db_path, backup_path)
            log.info(f"Backed up active db to {backup_path.name}")
            
        shutil.move(staging_path, db_path)
        log.info(f"Active database updated. Loaded {files_loaded} files, {total_loaded} rows.")
        return True
    else:
        log.warning("NSC active database was not replaced due to errors.")
        return False

def main():
    parser = argparse.ArgumentParser(description="NSC Data Importer")
    parser.add_argument("--config", help="Path to database.yaml")
    parser.add_argument("--dry-run", action="store_true", help="Discover files without importing")
    args = parser.parse_args()
    
    config = load_config(args.config)
    log_dir = Path(project_root) / config.get('logging', {}).get('log_dir', 'Validation/Database/logs')
    log_level = getattr(logging, config.get('logging', {}).get('level', 'INFO').upper(), logging.INFO)
    
    setup_logger("nsc_importer", log_dir=log_dir, level=log_level)
    
    try:
        run_import(config, dry_run=args.dry_run)
    except KeyboardInterrupt:
        logging.getLogger("nsc_importer").info("Import cancelled by user.")
        sys.exit(130)
    except Exception as e:
        logging.getLogger("nsc_importer").exception("Fatal error during import.")
        sys.exit(1)

if __name__ == "__main__":
    main()
